#!/usr/bin/env python3
"""Validate study-reader JSON and build a Word reader with python-docx.

Usage:
    python build_reader.py reader.json --check-only
    python build_reader.py reader.json --output reader.docx

Content validation uses only the standard library. The Word builder imports
python-docx lazily and reads appearance tokens from assets/word-style.json.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import tempfile
import unicodedata
from collections.abc import Mapping
from pathlib import Path
from types import SimpleNamespace


STYLE_PATH = Path(__file__).resolve().parents[1] / "assets" / "word-style.json"
FONT_KEYS = {"cjk", "english", "body", "heading"}


def _error(path: str, message: str) -> None:
    raise ValueError(f"{path}: {message}")


def _object(value, path, required, optional=()):
    if not isinstance(value, dict):
        _error(path, "must be an object")
    for key in required:
        if key not in value:
            _error(f"{path}.{key}", "required field is missing")
    unknown = set(value) - set(required) - set(optional)
    if unknown:
        _error(path, "unknown field(s): " + ", ".join(sorted(map(str, unknown))))
    return value


def _text(value, path, single_paragraph=False):
    if not isinstance(value, str) or not value.strip():
        _error(path, "must be a non-empty string")
    for char in value:
        code = ord(char)
        if ((code < 32 and char not in "\t\n\r") or
                0xD800 <= code <= 0xDFFF or code in (0xFFFE, 0xFFFF)):
            _error(path, f"contains an invalid XML character U+{code:04X}")
    if single_paragraph and any(c in value for c in "\r\n\u2028\u2029"):
        _error(path, "source sentence must be one continuous paragraph, without line breaks")
    return value


def _list(value, path, minimum=0):
    if not isinstance(value, list):
        _error(path, "must be an array")
    if len(value) < minimum:
        _error(path, f"must contain at least {minimum} item(s)")
    return value


def _strings(value, path, minimum=0):
    for i, item in enumerate(_list(value, path, minimum)):
        _text(item, f"{path}[{i}]")


def _unique(value, seen, path):
    if value in seen:
        _error(path, f"duplicate identifier {value!r}")
    seen.add(value)


def _source_index(data):
    return {
        sentence["id"]: sentence["text"]
        for paragraph in data["source"]
        for sentence in paragraph["sentences"]
    }


def _exercise_id(entry, index):
    return entry["id"] if len(entry["practice"]) == 1 else f"{entry['id']}.{index + 1}"


def _normalized_answer(answer):
    return " ".join(answer.split()).casefold()


def _source_refs(value, path, known_ids):
    _strings(value, path, 1)
    seen = set()
    for i, source_id in enumerate(value):
        rpath = f"{path}[{i}]"
        _unique(source_id, seen, rpath)
        if source_id not in known_ids:
            _error(rpath, f"unknown source reference {source_id!r}")


def validate(data) -> None:
    """Validate all content and references; raise ValueError with a JSON path.

    Sentence entries reference source sentence IDs. Evidence may reference a
    known source sentence or paragraph ID. Reviews require known, differing
    student and correct answers; provenance must still be checked by the author.
    No maximum count is imposed on selected sentences or their exercises.
    """
    _object(data, "$", ("meta", "source", "sentences", "vocabulary", "expressions"),
            ("reviews", "checks", "question_groups", "paraphrases"))
    meta = _object(data["meta"], "$.meta", ("title", "subtitle", "label"),
                   ("note", "answer_label", "include_review_sections"))
    for key in meta:
        if key == "include_review_sections":
            if not isinstance(meta[key], bool):
                _error("$.meta.include_review_sections", "must be a boolean")
        else:
            _text(meta[key], f"$.meta.{key}")

    paragraph_ids, sentence_ids = set(), set()
    for i, paragraph in enumerate(_list(data["source"], "$.source", 1)):
        path = f"$.source[{i}]"
        _object(paragraph, path, ("id", "sentences"))
        _text(paragraph["id"], path + ".id")
        _unique(paragraph["id"], paragraph_ids, path + ".id")
        for j, sentence in enumerate(_list(paragraph["sentences"], path + ".sentences", 1)):
            spath = f"{path}.sentences[{j}]"
            _object(sentence, spath, ("id", "text"))
            _text(sentence["id"], spath + ".id")
            _unique(sentence["id"], sentence_ids, spath + ".id")
            _text(sentence["text"], spath + ".text", single_paragraph=True)
    overlap = paragraph_ids & sentence_ids
    if overlap:
        _error("$.source", "paragraph and sentence IDs must be distinct: " + ", ".join(sorted(overlap)))

    selected_ids, exercise_ids = set(), set()
    for i, entry in enumerate(_list(data["sentences"], "$.sentences", 1)):
        path = f"$.sentences[{i}]"
        _object(entry, path, ("id", "title", "translation", "analysis", "practice"), ("depth",))
        depth = entry.get("depth", "full")
        if depth not in ("full", "brief"):
            _error(path + ".depth", "must be 'full' or 'brief'")
        for key in ("id", "title", "translation"):
            _text(entry[key], f"{path}.{key}")
        _unique(entry["id"], selected_ids, path + ".id")
        if entry["id"] not in sentence_ids:
            _error(path + ".id", f"unknown source sentence {entry['id']!r}")
        analysis = _object(entry["analysis"], path + ".analysis", ("core", "links"), ("focus",))
        _strings(analysis["core"], path + ".analysis.core", 1 if depth == "full" else 0)
        for j, link in enumerate(_list(analysis["links"], path + ".analysis.links", 1)):
            lpath = f"{path}.analysis.links[{j}]"
            _object(link, lpath, ("anchor", "relation", "detail"))
            for key in link:
                _text(link[key], f"{lpath}.{key}")
        focus = analysis.get("focus")
        if focus is not None:
            _object(focus, path + ".analysis.focus", ("label", "text"))
            for key in focus:
                _text(focus[key], f"{path}.analysis.focus.{key}")
        for j, exercise in enumerate(_list(entry["practice"], path + ".practice", 1 if depth == "full" else 0)):
            epath = f"{path}.practice[{j}]"
            _object(exercise, epath, ("text", "questions", "answer"))
            _text(exercise["text"], epath + ".text")
            _strings(exercise["questions"], epath + ".questions", 1)
            _text(exercise["answer"], epath + ".answer")
            exercise_ids.add(_exercise_id(entry, j))

    check_ids = set()
    for i, check in enumerate(_list(data.get("checks", []), "$.checks")):
        path = f"$.checks[{i}]"
        _object(check, path, ("id", "title", "text", "questions", "answer"), ("source_ids",))
        for key in ("id", "title", "text", "answer"):
            _text(check[key], f"{path}.{key}")
        _unique(check["id"], check_ids, path + ".id")
        if check["id"] in exercise_ids:
            _error(path + ".id", f"conflicts with exercise identifier {check['id']!r}")
        _strings(check["questions"], path + ".questions", 1)
        if "source_ids" in check:
            _strings(check["source_ids"], path + ".source_ids")
            seen_refs = set()
            for j, source_id in enumerate(check["source_ids"]):
                rpath = f"{path}.source_ids[{j}]"
                _unique(source_id, seen_refs, rpath)
                if source_id not in sentence_ids | paragraph_ids:
                    _error(rpath, f"unknown source reference {source_id!r}")

    question_answers, question_numbers = {}, set()
    for i, group in enumerate(_list(data.get("question_groups", []), "$.question_groups")):
        path = f"$.question_groups[{i}]"
        _object(group, path, ("title", "instruction", "items"))
        for key in ("title", "instruction"):
            _text(group[key], f"{path}.{key}")
        for j, item in enumerate(_list(group["items"], path + ".items", 1)):
            ipath = f"{path}.items[{j}]"
            _object(item, ipath, ("number", "prompt", "source_ids", "answer", "explanation", "action"))
            for key in ("number", "prompt", "answer", "explanation", "action"):
                _text(item[key], f"{ipath}.{key}")
            _unique(item["number"], question_numbers, ipath + ".number")
            _source_refs(item["source_ids"], ipath + ".source_ids", sentence_ids | paragraph_ids)
            question_answers[item["number"]] = _normalized_answer(item["answer"])

    review_numbers = set()
    for i, review in enumerate(_list(data.get("reviews", []), "$.reviews")):
        path = f"$.reviews[{i}]"
        _object(review, path,
                ("number", "prompt", "user_answer", "correct_answer", "evidence", "paraphrases", "why_selected_fails"),
                ("hypothesis", "check"))
        for key in ("number", "prompt", "user_answer", "correct_answer", "why_selected_fails", "hypothesis", "check"):
            if key in review:
                _text(review[key], f"{path}.{key}")
        _unique(review["number"], review_numbers, path + ".number")
        user_answer = _normalized_answer(review["user_answer"])
        correct_answer = _normalized_answer(review["correct_answer"])
        if review["number"] in question_answers and correct_answer != question_answers[review["number"]]:
            _error(path + ".correct_answer", "conflicts with question_groups answer for the same number")
        # Words such as 'unknown' or 'none' can be legitimate gap-fill answers.
        # Whether a submission was actually observed is an authoring judgment.
        if user_answer == correct_answer:
            _error(path, "reviews must describe verified errors; user_answer and correct_answer are identical")
        for j, evidence in enumerate(_list(review["evidence"], path + ".evidence", 1)):
            epath = f"{path}.evidence[{j}]"
            _object(evidence, epath, ("source_ids", "explanation"))
            _strings(evidence["source_ids"], epath + ".source_ids", 1)
            seen_refs = set()
            for k, source_id in enumerate(evidence["source_ids"]):
                rpath = f"{epath}.source_ids[{k}]"
                _unique(source_id, seen_refs, rpath)
                if source_id not in sentence_ids | paragraph_ids:
                    _error(rpath, f"unknown source reference {source_id!r}")
            _text(evidence["explanation"], epath + ".explanation")
        _strings(review["paraphrases"], path + ".paraphrases")

    for i, item in enumerate(_list(data.get("paraphrases", []), "$.paraphrases")):
        path = f"$.paraphrases[{i}]"
        _object(item, path, ("origin", "source_ids", "source_text", "target_text", "note"), ("question_numbers",))
        if item["origin"] not in ("original", "transfer"):
            _error(path + ".origin", "must be 'original' or 'transfer'")
        for key in ("source_text", "target_text", "note"):
            _text(item[key], f"{path}.{key}")
        _source_refs(item["source_ids"], path + ".source_ids", sentence_ids | paragraph_ids)
        numbers = item.get("question_numbers", [])
        _strings(numbers, path + ".question_numbers", 1 if item["origin"] == "original" else 0)
        if item["origin"] == "transfer" and numbers:
            _error(path + ".question_numbers", "transfer paraphrases must not claim original question numbers")
        for j, number in enumerate(numbers):
            if number not in question_numbers:
                _error(f"{path}.question_numbers[{j}]", f"unknown original question number {number!r}")

    for key, fields in (("vocabulary", ("term", "meaning", "context")),
                        ("expressions", ("phrase", "meaning", "example", "translation"))):
        for i, item in enumerate(_list(data[key], f"$.{key}")):
            path = f"$.{key}[{i}]"
            _object(item, path, fields)
            for field in fields:
                _text(item[field], f"{path}.{field}")


def _load_style(fonts=None):
    try:
        style = json.loads(STYLE_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"word-style.json: cannot read style tokens: {exc}") from exc
    groups = {
        "fonts": FONT_KEYS,
        "colors": {"ink", "source", "blue", "teal", "muted", "table_header", "table_alternate", "table_border"},
        "page": {"width", "height", "top", "bottom", "left", "right", "footer"},
        "size": {"title", "heading1", "heading2", "body", "translation", "small", "table", "footer"},
        "spacing": {"body_line", "translation_line", "paragraph_after", "analysis_after", "heading_before", "heading_after"},
    }
    if not isinstance(style, dict):
        _error("word-style.json", "must be an object")
    for group, keys in groups.items():
        if group not in style or not isinstance(style[group], dict):
            _error(f"word-style.json.{group}", "required token group is missing")
        for key in keys:
            path = f"word-style.json.{group}.{key}"
            if key not in style[group]:
                _error(path, "required token is missing")
            value = style[group][key]
            if group in ("fonts", "colors"):
                _text(value, path)
                if group == "colors" and not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                    _error(path, "must be a six-digit RGB hex color")
            else:
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                    _error(path, "must be a finite number")
                allow_zero = group == "page" and key not in ("width", "height")
                allow_zero |= group == "spacing" and key not in ("body_line", "translation_line")
                if value < 0 or (not allow_zero and value == 0):
                    _error(path, "must be positive" if not allow_zero else "must be non-negative")
    page = style["page"]
    if page["left"] + page["right"] >= page["width"] or page["top"] + page["bottom"] >= page["height"]:
        _error("word-style.json.page", "margins leave no usable page area")
    if fonts is not None:
        if not isinstance(fonts, Mapping):
            _error("fonts", "must be a mapping of font tokens")
        if set(fonts) - FONT_KEYS:
            _error("fonts", "unknown font token(s): " + ", ".join(sorted(set(fonts) - FONT_KEYS)))
        for key, value in fonts.items():
            _text(value, f"fonts.{key}")
            style["fonts"][key] = value
    return style


def _docx_api():
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("Building DOCX requires python-docx; validation with --check-only does not.") from exc
    return SimpleNamespace(**locals())


class _Reader:
    def __init__(self, data, style):
        self.data, self.style, self.api = data, style, _docx_api()
        self.doc = self.api.Document()
        self.fonts, self.colors = style["fonts"], style["colors"]
        self.size, self.spacing, self.page = style["size"], style["spacing"], style["page"]
        self.width = (self.page["width"] - self.page["left"] - self.page["right"]) * 72
        self.height = (self.page["height"] - self.page["top"] - self.page["bottom"]) * 72
        self._setup()

    def _font(self, target, token="body", size=None, bold=False, color=None):
        a = self.api
        target.font.name = self.fonts[token]
        target.font.size = a.Pt(size if size is not None else self.size["body"])
        target.font.bold = bold
        target.font.color.rgb = a.RGBColor.from_string(color or self.colors["ink"])
        rp = target.element.get_or_add_rPr()
        fonts = rp.find(a.qn("w:rFonts"))
        if fonts is None:
            fonts = a.OxmlElement("w:rFonts")
            rp.insert(0, fonts)
        for key in list(fonts.attrib):
            if "theme" in key.lower():
                del fonts.attrib[key]
        for name in ("ascii", "hAnsi", "cs"):
            fonts.set(a.qn("w:" + name), self.fonts[token])
        fonts.set(a.qn("w:eastAsia"), self.fonts["cjk"])
        # The default Word template can carry a blue rule in the Title style.
        for border in list(target.element.findall(".//" + a.qn("w:pBdr"))):
            border.getparent().remove(border)

    def _setup(self):
        a, doc = self.api, self.doc
        section = doc.sections[0]
        section.page_width, section.page_height = a.Inches(self.page["width"]), a.Inches(self.page["height"])
        for edge in ("top", "bottom", "left", "right"):
            setattr(section, edge + "_margin", a.Inches(self.page[edge]))
        section.footer_distance = a.Inches(self.page["footer"])
        for name in ("Normal", "Body Text"):
            self._font(doc.styles[name])
            fmt = doc.styles[name].paragraph_format
            fmt.line_spacing = a.Pt(self.spacing["body_line"])
            fmt.space_after = a.Pt(self.spacing["paragraph_after"])
            fmt.widow_control = True
        for name, token in (("Title", "title"), ("Heading 1", "heading1"), ("Heading 2", "heading2")):
            self._font(doc.styles[name], "heading", self.size[token], True, color="000000")
            fmt = doc.styles[name].paragraph_format
            fmt.line_spacing = 1.15
            fmt.space_before, fmt.space_after = a.Pt(self.spacing["heading_before"]), a.Pt(self.spacing["heading_after"])
            fmt.keep_with_next = True
        doc.styles["Title"].paragraph_format.space_before = a.Pt(0)
        doc.styles["Title"].paragraph_format.space_after = a.Pt(5)
        custom = (
            ("English", "english", "body", "source"),
            ("Chinese", "body", "translation", "muted"),
            ("Small", "body", "small", "muted"),
            ("Reader Subtitle", "english", "body", "ink"),
            ("Practice", "english", "body", "ink"),
            ("Structure", "body", "body", "blue"),
        )
        for name, font_token, size_token, color_token in custom:
            st = doc.styles.add_style(name, 1)
            st.base_style = doc.styles["Normal"]
            self._font(st, font_token, self.size[size_token], color=self.colors[color_token])
        doc.styles["Reader Subtitle"].font.italic = True
        doc.styles["Reader Subtitle"].font.color.rgb = a.RGBColor.from_string("000000")
        doc.styles["Chinese"].paragraph_format.line_spacing = a.Pt(self.spacing["translation_line"])
        doc.styles["Small"].paragraph_format.line_spacing = a.Pt(self.size["small"] * 1.5)
        self._font(doc.styles["Footer"], "heading", self.size["footer"], color=self.colors["muted"])
        footer = section.footer.paragraphs[0]
        footer.alignment = a.WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(self.data["meta"]["label"] + "  ·  ")
        field = a.OxmlElement("w:fldSimple")
        field.set(a.qn("w:instr"), "PAGE")
        footer._p.append(field)

    def p(self, text="", style=None, keep=False):
        paragraph = self.doc.add_paragraph(style=style)
        paragraph.add_run(text)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = keep
        return paragraph

    def h(self, text, level=1, new_page=False):
        paragraph = self.p(text, "Heading " + str(level), keep=True)
        paragraph.paragraph_format.page_break_before = new_page
        return paragraph

    def en(self, text, keep=False):
        paragraph = self.p(text, "English", keep=keep)
        # Source text is copied as a single run, never split or highlighted.
        paragraph.runs[0].bold = False
        return paragraph

    def lead(self, label, text, style=None, color="blue", keep=False):
        paragraph = self.p("", style, keep)
        run = paragraph.add_run(label + "  " if label else "")
        run.bold = True
        run.font.color.rgb = self.api.RGBColor.from_string(self.colors[color])
        paragraph.add_run(text)
        return paragraph

    def _estimated_height(self, paragraphs):
        """Conservative layout estimate, not a substitute for rendered QA."""
        height = 0.0
        for paragraph in paragraphs:
            style = paragraph.style
            size = style.font.size.pt if style.font.size else self.size["body"]
            fmt, inherited = paragraph.paragraph_format, style.paragraph_format
            line = fmt.line_spacing if fmt.line_spacing is not None else inherited.line_spacing
            line_height = line.pt if hasattr(line, "pt") else size * float(line or self.spacing["body_line"] / size)
            width = self.width - (fmt.left_indent.pt if fmt.left_indent else 0) - (fmt.right_indent.pt if fmt.right_indent else 0)
            lines = 0
            for text in paragraph.text.split("\n"):
                units = sum(1 if unicodedata.east_asian_width(c) in ("W", "F") else
                            0.3 if c.isspace() else 0.56 for c in text)
                lines += max(1, math.ceil(units * size / max(width, size)))
            before = fmt.space_before if fmt.space_before is not None else inherited.space_before
            after = fmt.space_after if fmt.space_after is not None else inherited.space_after
            height += lines * line_height + (before.pt if before else 0) + (after.pt if after else self.spacing["paragraph_after"])
        return height

    def keep_block(self, start, max_fraction=0.90):
        paragraphs = self.doc.paragraphs[start:]
        # Naturally flow several short entries onto a page. Only chain a block
        # that should fit on one page; unusually long entries remain splittable.
        if self._estimated_height(paragraphs) <= self.height * max_fraction:
            for paragraph in paragraphs[:-1]:
                paragraph.paragraph_format.keep_with_next = True
        if paragraphs:
            paragraphs[-1].paragraph_format.keep_with_next = False

    @staticmethod
    def exercise_id(entry, index):
        return _exercise_id(entry, index)

    def questions(self, questions):
        return self.p(" ".join(
            (chr(0x2460 + i) if i < 20 else f"{i + 1}.") + question
            for i, question in enumerate(questions)
        ))

    def sentence(self, entry, original):
        start = len(self.doc.paragraphs)
        self.h(entry["id"] + " " + entry["title"], level=2)
        self.en(original, keep=True)
        self.lead("译文", entry["translation"], "Chinese", "muted", keep=True)
        for i, text in enumerate(entry["analysis"]["core"]):
            p = self.lead("结构示意" if i == 0 else "", text, "Structure", keep=True)
            p.paragraph_format.space_after = self.api.Pt(self.spacing["analysis_after"])
        for link in entry["analysis"]["links"]:
            p = self.p()
            p.paragraph_format.left_indent = self.api.Inches(.12)
            p.paragraph_format.first_line_indent = self.api.Inches(-.12)
            p.paragraph_format.space_after = self.api.Pt(self.spacing["analysis_after"])
            run = p.add_run("• " + link["anchor"])
            run.bold = True
            run.font.color.rgb = self.api.RGBColor.from_string(self.colors["blue"])
            p.add_run(" → " + link["relation"] + "。" + link["detail"])
        focus = entry["analysis"].get("focus")
        if focus:
            self.lead(focus["label"], focus["text"])
        # Even a large entry keeps its final explanation next to the exercise.
        if entry["practice"]:
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        for i, exercise in enumerate(entry["practice"]):
            p = self.p("", "Practice", keep=True)
            label = p.add_run("练习 " + self.exercise_id(entry, i) + "  ")
            self._font(label, "body", bold=True, color=self.colors["teal"])
            p.add_run(exercise["text"])
            self.questions(exercise["questions"])
        # Keep only compact entries whole. Longer explanations flow naturally;
        # otherwise most passages would inflate to one sentence per page.
        self.keep_block(start, max_fraction=0.45)

    def question_groups(self):
        if not self.data.get("question_groups"):
            return
        self.h("原题考点与快速核对", new_page=True)
        for group in self.data["question_groups"]:
            self.h(group["title"], level=2)
            self.p(group["instruction"], "Small", keep=True)
            self.table(["题目", "原文依据与答案", "核对后怎样推进"], [
                [f"第 {item['number']} 题\n{item['prompt']}",
                 "原文 " + "、".join(item["source_ids"]) + f"\n答案 {item['answer']}\n{item['explanation']}",
                 item["action"]]
                for item in group["items"]
            ], width_ratios=(.32, .44, .24))

    def reviews(self):
        if not self.data.get("reviews"):
            return
        self.h("错题复盘")
        mapped_numbers = {item["number"] for group in self.data.get("question_groups", []) for item in group["items"]}
        for review in self.data["reviews"]:
            start = len(self.doc.paragraphs)
            answer_label = self.data["meta"].get("answer_label", "你的答案")
            self.h(f"第 {review['number']} 题  {answer_label} {review['user_answer']}  正确答案 {review['correct_answer']}", level=2)
            if review["number"] not in mapped_numbers:
                self.en(review["prompt"], keep=True)
                for evidence in review["evidence"]:
                    self.lead("证据 " + "、".join(evidence["source_ids"]), evidence["explanation"])
                if review["paraphrases"]:
                    self.lead("题干改写", "；".join(review["paraphrases"]))
            self.lead("答案对照", review["why_selected_fails"])
            if "hypothesis" in review:
                self.lead("可能的卡点", review["hypothesis"])
            if "check" in review:
                self.lead("核对一下", review["check"])
            self.keep_block(start)

    def checks(self):
        if not self.data.get("checks"):
            return
        self.h("换材料再判断")
        for check in self.data["checks"]:
            start = len(self.doc.paragraphs)
            self.h(check["id"] + " " + check["title"], level=2)
            # Blank lines mark natural paragraph boundaries in a check passage.
            # Internal source references and answers never appear beside it.
            text = check["text"].replace("\r\n", "\n").replace("\r", "\n")
            text = text.replace("\u2029", "\n\n").replace("\u2028", "\n")
            for paragraph in re.split(r"\n[ \t]*\n+", text):
                if paragraph.strip():
                    self.en(paragraph.strip(), keep=True)
            self.questions(check["questions"])
            self.keep_block(start)

    def paraphrases(self):
        if not self.data.get("paraphrases"):
            return
        self.h("考点表达与同义改写", new_page=True)
        for origin, title, target_header in (("original", "原题改写", "原题改写"),
                                              ("transfer", "原创迁移说法", "原创迁移说法")):
            items = [item for item in self.data["paraphrases"] if item["origin"] == origin]
            if not items:
                continue
            self.h(title, level=2)
            rows = []
            for item in items:
                references = "原文 " + "、".join(item["source_ids"])
                if origin == "original":
                    references += "；原题 " + "、".join(item["question_numbers"])
                rows.append([item["source_text"], item["target_text"], item["note"] + "\n" + references])
            self.table(["原文表达", target_header, "关系与使用边界"], rows, width_ratios=(.29, .32, .39))

    def table(self, headers, rows, width_ratios=None):
        a = self.api
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        widths = [self.width / 72 * ratio for ratio in (width_ratios or (.21, .24, .55))]
        for column, width in zip(table.columns, widths):
            column.width = a.Inches(width)
        borders = a.OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = a.OxmlElement("w:" + side)
            for key, value in (("val", "single"), ("sz", "4"), ("color", self.colors["table_border"])):
                edge.set(a.qn("w:" + key), value)
            borders.append(edge)
        table._tbl.tblPr.append(borders)
        margins = a.OxmlElement("w:tblCellMar")
        for edge, value in (("top", "65"), ("bottom", "65"), ("left", "120"), ("right", "120")):
            element = a.OxmlElement("w:" + edge)
            element.set(a.qn("w:w"), value)
            element.set(a.qn("w:type"), "dxa")
            margins.append(element)
        table._tbl.tblPr.append(margins)
        for i, row in enumerate([headers] + rows):
            cells = table.rows[0].cells if i == 0 else table.add_row().cells
            trpr = cells[0]._tc.getparent().get_or_add_trPr()
            trpr.append(a.OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(a.OxmlElement("w:tblHeader"))
            for j, (cell, text) in enumerate(zip(cells, row)):
                cell.width = a.Inches(widths[j])
                cell.vertical_alignment = a.WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = p.paragraph_format.space_after = a.Pt(3)
                p.paragraph_format.line_spacing = a.Pt(self.size["table"] * 1.43)
                self._font(p.add_run(text), size=self.size["table"], bold=(i == 0 or j == 0),
                           color="FFFFFF" if i == 0 else self.colors["blue"] if j == 0 else self.colors["ink"])
                shade = a.OxmlElement("w:shd")
                shade.set(a.qn("w:fill"), self.colors["table_header"] if i == 0 else
                          self.colors["table_alternate"] if i % 2 else "FFFFFF")
                cell._tc.get_or_add_tcPr().append(shade)

    def render(self):
        meta = self.data["meta"]
        include_reviews = meta.get("include_review_sections", False)
        printed_checks = self.data.get("checks", []) if include_reviews else []
        self.p(meta["title"], "Title", keep=True)
        self.p(meta["subtitle"], "Reader Subtitle", keep="note" in meta)
        if "note" in meta:
            self.p(meta["note"], "Small")
        source = _source_index(self.data)
        positions = {source_id: i for i, source_id in enumerate(source)}
        entries = sorted(self.data["sentences"], key=lambda entry: positions[entry["id"]])
        for entry in entries:
            self.sentence(entry, source[entry["id"]])
        if include_reviews:
            self.question_groups()
            self.reviews()
            self.checks()
        self.paraphrases()
        if self.data["vocabulary"]:
            self.h("阅读词汇与常见改写", new_page=True)
            self.table(["词或词组", "本篇意思", "放回原文理解"],
                       [[v["term"], v["meaning"], v["context"]] for v in self.data["vocabulary"]])
        if self.data["expressions"]:
            self.h("常用表达", new_page=True)
            for expression in self.data["expressions"]:
                start = len(self.doc.paragraphs)
                heading = self.h(expression["phrase"], level=2)
                meaning = heading.add_run("  " + expression["meaning"])
                self._font(meaning, size=self.size["translation"], color=self.colors["muted"])
                self.en(expression["example"], keep=True)
                self.p(expression["translation"], "Chinese")
                self.keep_block(start)
        if any(entry["practice"] for entry in entries) or printed_checks:
            self.h("练习参考答案", new_page=True)
        for entry in entries:
            if not entry["practice"]:
                continue
            start = len(self.doc.paragraphs)
            self.h(entry["id"] + " " + entry["title"], level=2)
            for i, exercise in enumerate(entry["practice"]):
                if len(entry["practice"]) > 1:
                    self.lead("练习 " + self.exercise_id(entry, i), exercise["answer"], color="teal")
                else:
                    self.p(exercise["answer"])
            self.keep_block(start)
        for check in printed_checks:
            start = len(self.doc.paragraphs)
            self.h("检查 " + check["id"] + " " + check["title"], level=2)
            self.p(check["answer"])
            self.keep_block(start)
        self.doc.core_properties.title = meta["title"]
        self.doc.core_properties.subject = meta["subtitle"]
        self.doc.core_properties.author = ""
        return self.doc


def build(data, output, fonts=None) -> Path:
    """Validate and atomically write a DOCX; return its path.

    ``fonts`` optionally overrides cjk, english, body, or heading tokens.
    Validation and document assembly complete before the output is touched.
    """
    validate(data)
    style = _load_style(fonts)
    output = Path(output)
    if output.suffix.lower() != ".docx":
        _error("output", "must have a .docx extension")
    document = _Reader(data, style).render()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix="." + output.stem + "-", suffix=".docx", dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
        document.save(temporary)
        os.replace(temporary, output)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return output


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Reader JSON file (UTF-8)")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--output", type=Path, help="Destination .docx file")
    mode.add_argument("--check-only", action="store_true", help="Validate content without building a document")
    for token in sorted(FONT_KEYS):
        parser.add_argument("--" + token + "-font", help=f"Override the {token} font token")
    args = parser.parse_args(argv)
    try:
        data = json.loads(args.input.read_text(encoding="utf-8"))
        if args.check_only:
            validate(data)
            print("Validation passed.")
        else:
            fonts = {token: getattr(args, token + "_font") for token in FONT_KEYS if getattr(args, token + "_font") is not None}
            print(build(data, args.output, fonts=fonts))
    except (ValueError, OSError, RuntimeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
