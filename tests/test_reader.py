"""Observable checks for source fidelity, practice separation and invalid input."""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('reader', ROOT / 'scripts/build_reader.py')
reader = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(reader)


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads((ROOT / 'tests/fixtures/reader.json').read_text(encoding='utf-8'))

    def test_originals_remain_continuous_and_answers_follow_all_exercises(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'reader.docx'
            reader.build(self.data, out)
            doc = Document(out)
            paragraphs = [p.text for p in doc.paragraphs]
            self.assertNotIn('原题考点与快速核对', paragraphs)
            self.assertNotIn('考点表达与同义改写', paragraphs)
            for expression in self.data['expressions']:
                self.assertIn(expression['phrase'] + '  ' + expression['meaning'], paragraphs)
            source = {s['id']: s['text'] for p in self.data['source'] for s in p['sentences']}
            for item in self.data['sentences']:
                found = [p for p in doc.paragraphs if p.text == source[item['id']]]
                self.assertEqual(len(found), 1)
                self.assertFalse(any(run.bold for run in found[0].runs))
            exercise_positions, answer_positions = [], []
            for item in self.data['sentences']:
                for practice in item['practice']:
                    exercise_positions.append(next(i for i, p in enumerate(paragraphs) if practice['text'] in p))
                    answer_positions.append(next(i for i, p in enumerate(paragraphs) if practice['answer'] in p))
            self.assertGreater(min(answer_positions), max(exercise_positions))

    def test_selection_count_is_not_fixed_and_review_evidence_can_be_unselected(self):
        self.data['sentences'] = self.data['sentences'][:1]
        # The review still cites C2, which is in the source but not selected above.
        reader.validate(self.data)
        with tempfile.TemporaryDirectory() as temp:
            reader.build(self.data, Path(temp) / 'one-entry.docx')

    def test_missing_exercise_unknown_evidence_and_duplicate_source_are_rejected(self):
        bad = copy.deepcopy(self.data)
        bad['sentences'][0]['practice'] = []
        with self.assertRaises(ValueError):
            reader.validate(bad)
        bad = copy.deepcopy(self.data)
        bad['reviews'][0]['evidence'][0]['source_ids'] = ['Z99']
        with self.assertRaises(ValueError):
            reader.validate(bad)
        bad = copy.deepcopy(self.data)
        bad['source'][1]['sentences'][0]['id'] = bad['source'][0]['sentences'][0]['id']
        with self.assertRaises(ValueError):
            reader.validate(bad)

    def test_invalid_input_does_not_replace_an_existing_output(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            out = folder / 'existing.docx'
            out.write_bytes(b'Existing user file')
            bad = folder / 'invalid.json'
            bad.write_text('{}', encoding='utf-8')
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/build_reader.py'), str(bad), '--output', str(out)], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(out.read_bytes(), b'Existing user file')

    def test_default_and_explicit_full_entries_still_require_core_and_practice(self):
        for depth in (None, 'full'):
            for missing in ('core', 'practice'):
                with self.subTest(depth=depth, missing=missing):
                    bad = copy.deepcopy(self.data)
                    entry = bad['sentences'][0]
                    if depth is not None:
                        entry['depth'] = depth
                    if missing == 'core':
                        entry['analysis']['core'] = []
                    else:
                        entry['practice'] = []
                    with self.assertRaisesRegex(ValueError, missing):
                        reader.validate(bad)
        self.data['sentences'][0]['depth'] = 'full'
        reader.validate(self.data)

    def test_brief_entry_retains_source_translation_and_links_without_empty_sections(self):
        self.data['sentences'] = self.data['sentences'][:1]
        entry = self.data['sentences'][0]
        entry['depth'] = 'brief'
        entry['analysis']['core'] = []
        entry['practice'] = []
        reader.validate(self.data)
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'brief.docx'
            reader.build(self.data, out)
            doc = Document(out)
            texts = [p.text for p in doc.paragraphs]
            original = next(s['text'] for p in self.data['source'] for s in p['sentences'] if s['id'] == entry['id'])
            source_paragraphs = [p for p in doc.paragraphs if p.text == original]
            self.assertEqual(len(source_paragraphs), 1)
            self.assertEqual(len(source_paragraphs[0].runs), 1)
            self.assertFalse(source_paragraphs[0].runs[0].bold)
            self.assertTrue(any(entry['translation'] in p for p in texts))
            for link in entry['analysis']['links']:
                self.assertTrue(any(link['anchor'] in p and link['detail'] in p for p in texts))
            self.assertFalse(any(p.startswith('结构示意') for p in texts))
            self.assertFalse(any(p.startswith('练习 ' + entry['id']) for p in texts))
            self.assertNotIn('练习参考答案', texts)
            self.assertFalse(any('brief' in p or 'depth' in p for p in texts))

    def test_brief_still_requires_links_and_depth_values_are_checked(self):
        entry = self.data['sentences'][0]
        entry['depth'] = 'brief'
        # Brief entries may still include a core and exercises when useful.
        reader.validate(self.data)
        entry['analysis']['links'] = []
        with self.assertRaisesRegex(ValueError, r'analysis\.links'):
            reader.validate(self.data)
        entry['depth'] = 'mastered'
        with self.assertRaisesRegex(ValueError, r'\.depth'):
            reader.validate(self.data)

    @staticmethod
    def check_item(identifier='M1'):
        return {
            'id': identifier,
            'title': '换一个场景判断',
            'text': 'The school planned to remove the benches.\n\nAfter pupils explained how they used them, the school kept the benches.',
            'questions': ['学校最后采取了什么措施？', '什么信息改变了原来的决定？'],
            'answer': 'CHECK ANSWER ' + identifier + '：保留长椅；学生解释了长椅的用途。',
            'source_ids': ['A', 'A1'],
        }

    def test_checks_follow_reviews_and_all_answers_stay_in_the_final_appendix(self):
        self.data['meta']['include_review_sections'] = True
        self.data['checks'] = [self.check_item('M2'), self.check_item('M1')]
        self.data['checks'][1]['text'] = 'The club moved its meetings online but kept the usual starting time.'
        del self.data['checks'][1]['source_ids']
        self.data['sentences'][0]['depth'] = 'brief'
        self.data['sentences'][0]['analysis']['core'] = []
        self.data['sentences'][0]['practice'] = []
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'checks.docx'
            reader.build(self.data, out)
            doc = Document(out)
            texts = [p.text for p in doc.paragraphs]
            checks_start = texts.index('换材料再判断')
            vocabulary_start = texts.index('阅读词汇与常见改写')
            answers_start = texts.index('练习参考答案')
            review_end = next(i for i, text in enumerate(texts) if self.data['reviews'][-1]['check'] in text)
            self.assertLess(review_end, checks_start)
            self.assertLess(checks_start, vocabulary_start)
            self.assertLess(vocabulary_start, answers_start)
            for check in self.data['checks']:
                for paragraph in check['text'].split('\n\n'):
                    matches = [p for p in doc.paragraphs if p.text == paragraph]
                    self.assertEqual(len(matches), 1)
                    self.assertFalse(any(run.bold for run in matches[0].runs))
                    self.assertGreater(texts.index(paragraph), checks_start)
                    self.assertLess(texts.index(paragraph), vocabulary_start)
                self.assertNotIn(check['answer'], '\n'.join(texts[:answers_start]))
                answer_heading = '检查 ' + check['id'] + ' ' + check['title']
                self.assertGreater(texts.index(answer_heading), answers_start)
                self.assertEqual(texts[texts.index(answer_heading) + 1], check['answer'])
            sentence_answers = [practice['answer'] for entry in self.data['sentences'] for practice in entry['practice']]
            for answer in sentence_answers:
                self.assertGreater(texts.index(answer), answers_start)
                self.assertLess(texts.index(answer), texts.index(self.data['checks'][0]['answer']))
            self.assertLess(texts.index(self.data['checks'][0]['answer']), texts.index(self.data['checks'][1]['answer']))
            # Association metadata must not expose a source clue beside the check.
            check_body = '\n'.join(texts[checks_start:vocabulary_start])
            self.assertNotIn('A1', check_body)
            brief_heading = self.data['sentences'][0]['id'] + ' ' + self.data['sentences'][0]['title']
            self.assertNotIn(brief_heading, texts[answers_start:])

    def test_checks_reject_duplicate_ids_exercise_collisions_and_unknown_references(self):
        cases = ('duplicate', 'single-exercise', 'multiple-exercises', 'unknown-source', 'empty-question')
        for case in cases:
            with self.subTest(case=case):
                bad = copy.deepcopy(self.data)
                bad['checks'] = [self.check_item()]
                check = bad['checks'][0]
                if case == 'duplicate':
                    bad['checks'].append(copy.deepcopy(check))
                elif case == 'single-exercise':
                    check['id'] = bad['sentences'][0]['id']
                elif case == 'multiple-exercises':
                    entry = bad['sentences'][0]
                    entry['practice'].append(copy.deepcopy(entry['practice'][0]))
                    check['id'] = entry['id'] + '.2'
                elif case == 'unknown-source':
                    check['source_ids'] = ['Z99']
                else:
                    check['questions'] = [' ']
                with self.assertRaisesRegex(ValueError, r'\$\.checks\['):
                    reader.validate(bad)

    def test_check_only_answers_render_when_all_sentence_entries_are_brief(self):
        self.data['meta']['include_review_sections'] = True
        self.data['sentences'] = self.data['sentences'][:1]
        entry = self.data['sentences'][0]
        entry['depth'] = 'brief'
        entry['analysis']['core'] = []
        entry['practice'] = []
        self.data['checks'] = [self.check_item()]
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'check-only-answers.docx'
            reader.build(self.data, out)
            texts = [p.text for p in Document(out).paragraphs]
            answer_start = texts.index('练习参考答案')
            check = self.data['checks'][0]
            self.assertEqual(texts[answer_start + 1], '检查 ' + check['id'] + ' ' + check['title'])
            self.assertEqual(texts[-1], check['answer'])

    def test_invalid_check_does_not_replace_existing_output(self):
        self.data['checks'] = [self.check_item(self.data['sentences'][0]['id'])]
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'existing.docx'
            out.write_bytes(b'Keep existing reader')
            with self.assertRaisesRegex(ValueError, 'conflicts with exercise'):
                reader.build(self.data, out)
            self.assertEqual(out.read_bytes(), b'Keep existing reader')

    def question_group(self):
        # Reuse the fictional demo question; no real test-bank material is added.
        review = self.data['reviews'][0]
        return {
            'title': 'Paragraph matching',
            'instruction': 'Which paragraph contains the following information?',
            'items': [{
                'number': review['number'],
                'prompt': review['prompt'],
                'source_ids': ['C', 'C2'],
                'answer': review['correct_answer'],
                'explanation': review['evidence'][0]['explanation'],
                'action': '核对调整决定与保留家具两个条件，再继续下一题。',
            }],
        }

    @staticmethod
    def paraphrase(origin='original'):
        item = {
            'origin': origin,
            'source_ids': ['C2'],
            'source_text': 'the existing chairs',
            'target_text': 'furniture the library already owned',
            'note': '原文具体指椅子，不能扩展为所有家具都未更换。',
        }
        if origin == 'original':
            item['question_numbers'] = ['2']
        else:
            item['target_text'] = 'chairs the library already had'
        return item

    def test_question_group_validation_and_normalized_answer_consistency(self):
        self.data['question_groups'] = [self.question_group()]
        self.data['question_groups'][0]['items'][0]['answer'] = ' c \n'
        reader.validate(self.data)
        cases = ('duplicate-number', 'unknown-source', 'empty-source', 'empty-group',
                 'empty-instruction', 'empty-prompt', 'nonstring-number', 'answer-conflict')
        for case in cases:
            with self.subTest(case=case):
                bad = copy.deepcopy(self.data)
                group = bad['question_groups'][0]
                item = group['items'][0]
                if case == 'duplicate-number':
                    bad['question_groups'].append(copy.deepcopy(group))
                elif case == 'unknown-source':
                    item['source_ids'] = ['Z99']
                elif case == 'empty-source':
                    item['source_ids'] = []
                elif case == 'empty-group':
                    group['items'] = []
                elif case == 'empty-instruction':
                    group['instruction'] = ' '
                elif case == 'empty-prompt':
                    item['prompt'] = ''
                elif case == 'nonstring-number':
                    item['number'] = 2
                else:
                    item['answer'] = 'A'
                with self.assertRaisesRegex(ValueError, r'\$\.(question_groups|reviews)\['):
                    reader.validate(bad)

    def test_question_groups_do_not_require_reviews_or_selected_evidence(self):
        self.data['meta']['include_review_sections'] = True
        self.data['question_groups'] = [self.question_group()]
        self.data['sentences'] = self.data['sentences'][:1]
        del self.data['reviews']
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'questions-without-reviews.docx'
            reader.build(self.data, out)
            texts = [p.text for p in Document(out).paragraphs]
            self.assertIn('原题考点与快速核对', texts)
            self.assertNotIn('错题复盘', texts)

    def test_paraphrases_reject_false_question_attribution_and_bad_references(self):
        self.data['question_groups'] = [self.question_group()]
        self.data['paraphrases'] = [self.paraphrase(), self.paraphrase('transfer')]
        reader.validate(self.data)
        self.data['paraphrases'][1]['question_numbers'] = []
        reader.validate(self.data)
        for case in ('missing-number', 'empty-number', 'unknown-number', 'transfer-number',
                     'unknown-source', 'empty-source', 'empty-text', 'bad-origin'):
            with self.subTest(case=case):
                bad = copy.deepcopy(self.data)
                item = bad['paraphrases'][0]
                if case == 'missing-number':
                    del item['question_numbers']
                elif case == 'empty-number':
                    item['question_numbers'] = []
                elif case == 'unknown-number':
                    item['question_numbers'] = ['99']
                elif case == 'transfer-number':
                    item['origin'] = 'transfer'
                elif case == 'unknown-source':
                    item['source_ids'] = ['Z99']
                elif case == 'empty-source':
                    item['source_ids'] = []
                elif case == 'empty-text':
                    item['target_text'] = ' '
                else:
                    item['origin'] = 'inferred'
                with self.assertRaisesRegex(ValueError, r'\$\.paraphrases\['):
                    reader.validate(bad)

    def test_new_tables_follow_section_order_and_keep_original_and_transfer_distinct(self):
        self.data['meta']['include_review_sections'] = True
        self.data['question_groups'] = [self.question_group()]
        self.data['checks'] = [self.check_item()]
        self.data['paraphrases'] = [self.paraphrase('transfer'), self.paraphrase()]
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'question-tables.docx'
            reader.build(self.data, out)
            doc = Document(out)
            body = list(doc.element.body)
            sections = ('原题考点与快速核对', '错题复盘', '换材料再判断',
                        '考点表达与同义改写', '阅读词汇与常见改写', '常用表达', '练习参考答案')
            headings = {p.text: p for p in doc.paragraphs if p.text in sections}
            positions = [body.index(headings[text]._p) for text in sections]
            self.assertEqual(positions, sorted(positions))
            last_exercise = self.data['sentences'][-1]['practice'][-1]['text']
            last_exercise_p = next(p for p in doc.paragraphs if last_exercise in p.text)
            self.assertLess(body.index(last_exercise_p._p), positions[0])
            for title in ('原题考点与快速核对', '考点表达与同义改写'):
                self.assertTrue(headings[title].paragraph_format.page_break_before)
            self.assertIn(self.data['question_groups'][0]['instruction'], [p.text for p in doc.paragraphs])
            question_table, original_table, transfer_table = doc.tables[:3]
            self.assertEqual([c.text for c in question_table.rows[0].cells],
                             ['题目', '原文依据与答案', '核对后怎样推进'])
            item = self.data['question_groups'][0]['items'][0]
            self.assertIn(item['prompt'], question_table.cell(1, 0).text)
            for value in (item['answer'], item['explanation'], 'C、C2'):
                self.assertIn(value, question_table.cell(1, 1).text)
            self.assertEqual(item['action'], question_table.cell(1, 2).text)
            self.assertEqual(original_table.cell(0, 1).text, '原题改写')
            self.assertIn('原题 2', original_table.cell(1, 2).text)
            self.assertEqual(transfer_table.cell(0, 1).text, '原创迁移说法')
            self.assertNotIn('原题', transfer_table.cell(1, 2).text)
            for table, start, end in ((question_table, positions[0], positions[1]),
                                      (original_table, positions[3], positions[4]),
                                      (transfer_table, positions[3], positions[4])):
                self.assertLess(start, body.index(table._tbl))
                self.assertLess(body.index(table._tbl), end)
                self.assertTrue(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
                for row in table.rows:
                    self.assertTrue(row._tr.xpath('./w:trPr/w:cantSplit'))
                    self.assertFalse(row._tr.xpath('./w:trPr/w:trHeight[@w:hRule="exact"]'))
                    for cell in row.cells:
                        self.assertEqual(cell.paragraphs[0].runs[0].font.size.pt, 10.5)
            answers_start = next(i for i, p in enumerate(doc.paragraphs) if p.text == '练习参考答案')
            before_answers = '\n'.join(p.text for p in doc.paragraphs[:answers_start])
            self.assertNotIn(self.data['checks'][0]['answer'], before_answers)

    def test_mapped_review_omits_repetition_but_keeps_error_reason_and_diagnosis(self):
        self.data['meta']['include_review_sections'] = True
        self.data['question_groups'] = [self.question_group()]
        review = self.data['reviews'][0]
        unmapped = copy.deepcopy(review)
        unmapped['number'] = 'unmapped-fixture'
        unmapped['prompt'] = 'Unmapped fixture prompt'
        unmapped['evidence'][0]['explanation'] = 'Unmapped fixture evidence'
        unmapped['paraphrases'] = ['Unmapped fixture paraphrase']
        self.data['reviews'].append(unmapped)
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'mapped-review.docx'
            reader.build(self.data, out)
            texts = [p.text for p in Document(out).paragraphs]
            review_text = '\n'.join(texts[texts.index('错题复盘'):])
            for field in ('prompt',):
                self.assertNotIn(review[field], review_text)
            self.assertNotIn(review['evidence'][0]['explanation'], review_text)
            for paraphrase in review['paraphrases']:
                self.assertNotIn(paraphrase, review_text)
            for field in ('why_selected_fails', 'hypothesis', 'check'):
                self.assertIn(review[field], review_text)
            self.assertIn('虚构作答 B  正确答案 C', review_text)
            self.assertIn(unmapped['prompt'], review_text)
            self.assertIn(unmapped['evidence'][0]['explanation'], review_text)
            self.assertIn(unmapped['paraphrases'][0], review_text)


    def test_default_omits_review_sections_and_their_answers_but_keeps_learning_content(self):
        self.data['question_groups'] = [self.question_group()]
        self.data['checks'] = [self.check_item()]
        self.data['paraphrases'] = [self.paraphrase()]
        for explicit in (None, False):
            with self.subTest(explicit=explicit), tempfile.TemporaryDirectory() as temp:
                if explicit is not None:
                    self.data['meta']['include_review_sections'] = explicit
                out = Path(temp) / 'compact.docx'
                reader.build(self.data, out)
                doc = Document(out)
                paragraphs = [p.text for p in doc.paragraphs]
                all_text = '\n'.join(paragraphs + [c.text for t in doc.tables for row in t.rows for c in row.cells])
                for review in self.data['reviews']:
                    self.assertNotIn(review['why_selected_fails'], all_text)
                check = self.data['checks'][0]
                self.assertNotIn(check['answer'], all_text)
                self.assertNotIn(check['title'], all_text)
                self.assertNotIn(check['questions'][0], all_text)
                self.assertNotIn(self.data['question_groups'][0]['items'][0]['action'], all_text)
                for heading in ('原题考点与快速核对', '错题复盘', '换材料再判断'):
                    self.assertNotIn(heading, paragraphs)
                self.assertIn('考点表达与同义改写', paragraphs)
                self.assertIn(self.data['paraphrases'][0]['target_text'], all_text)
                self.assertIn('原题 2', all_text)
                for entry in self.data['sentences']:
                    for exercise in entry['practice']:
                        self.assertIn(exercise['text'], all_text)
                        self.assertIn(exercise['answer'], all_text)

    def test_hidden_checks_do_not_create_an_empty_answer_appendix(self):
        self.data['sentences'] = self.data['sentences'][:1]
        entry = self.data['sentences'][0]
        entry.update(depth='brief', practice=[])
        entry['analysis']['core'] = []
        self.data['checks'] = [self.check_item()]
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'no-answers.docx'
            reader.build(self.data, out)
            self.assertNotIn('练习参考答案', [p.text for p in Document(out).paragraphs])

    def test_review_section_option_rejects_non_boolean_values(self):
        for value in ('false', 0, 1, None, []):
            with self.subTest(value=value):
                self.data['meta']['include_review_sections'] = value
                with self.assertRaisesRegex(ValueError, 'include_review_sections'):
                    reader.validate(self.data)


if __name__ == '__main__':
    unittest.main()
